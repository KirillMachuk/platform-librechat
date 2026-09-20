#!/usr/bin/env python3
"""Exercise supported XLSX authoring cases against real LibreOffice output."""

from __future__ import annotations

import argparse
import copy
import hashlib
import json
import math
import subprocess
import sys
import tempfile
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.chart import BarChart, LineChart
from PIL import Image


ROOT = Path(__file__).resolve().parents[2]
BUILDER = ROOT / "skill/xlsx/scripts/build_spreadsheet.py"
FIXTURE = ROOT / "bench/spreadsheets/fixtures/ru_service_cost.json"


def _cases() -> dict[str, tuple[dict, dict]]:
    base = json.loads(FIXTURE.read_text(encoding="utf-8"))
    cases: dict[str, tuple[dict, dict]] = {}

    planned = copy.deepcopy(base)
    cases["planned_cost"] = (planned, {
        "sheets": ["Summary", "Данные"], "formulas": {"Данные!D4": "=B4*C4"},
        "values": {"Данные!D4": 72000, "Summary!B3": 378700, "Summary!B4": 6},
        "mutation": {"input": ("Данные!B4", 5),
                     "values": {"Данные!D4": 90000, "Summary!B3": 396700}},
        "chart": True, "chart_value_column": "D",
        "text": ["Показатель", "Стоимость по услугам", "Источник:"],
    })

    single = copy.deepcopy(base)
    single["summary"] = []
    single["chart"] = None
    cases["single_sheet"] = (single, {
        "sheets": ["Данные"], "formulas": {"Данные!D4": "=B4*C4"},
        "values": {"Данные!D4": 72000}, "chart": False,
        "text": ["Плановая стоимость услуг", "Стоимость, ₽"],
    })

    sourced = copy.deepcopy(base)
    sourced["sources"].append({"label": "Методика", "location": "Демонстрационное описание расчёта"})
    cases["source_trace"] = (sourced, {
        "sheets": ["Summary", "Данные", "Sources"],
        "formulas": {"Данные!D4": "=B4*C4"},
        "values": {"Данные!D4": 72000, "Summary!B3": 378700},
        "chart": True, "text": ["Источники", "Демонстрационное описание расчёта"],
    })

    dated = copy.deepcopy(base)
    dated["table"]["columns"].insert(1, {"key": "date", "header": "Дата", "type": "date"})
    for index, row in enumerate(dated["table"]["rows"], start=1):
        row["date"] = f"2026-09-{index:02d}"
    dated["chart"] = None
    cases["typed_dates"] = (dated, {
        "sheets": ["Summary", "Данные"], "formulas": {"Данные!E4": "=C4*D4"},
        "values": {"Данные!E4": 72000, "Summary!B3": 378700},
        "chart": False, "date": ("Данные!B4", datetime(2026, 9, 1)),
        "text": ["Дата", "Итого, ₽"],
    })

    discounted = copy.deepcopy(base)
    discounted["table"]["columns"].append(
        {"key": "discount", "header": "Скидка", "type": "percent", "minimum": 0, "maximum": 1}
    )
    for row in discounted["table"]["rows"]:
        row["discount"] = 0.1
    discounted["table"]["calculatedColumns"].extend([
        {
            "key": "discount_amount", "header": "Сумма скидки, ₽",
            "operation": "multiply", "inputs": ["cost", "discount"],
        },
        {
            "key": "net_cost", "header": "После скидки, ₽",
            "operation": "subtract", "inputs": ["cost", "discount_amount"],
        },
    ])
    discounted["summary"][0]["column"] = "net_cost"
    discounted["chart"]["value"] = "net_cost"
    cases["scenario_discount"] = (discounted, {
        "sheets": ["Summary", "Данные"],
        "formulas": {"Данные!E4": "=B4*C4", "Данные!F4": "=E4*D4", "Данные!G4": "=E4-F4"},
        "values": {"Данные!E4": 72000, "Данные!F4": 7200, "Данные!G4": 64800, "Summary!B3": 340830},
        "mutation": {"input": ("Данные!D4", 0.2),
                     "values": {"Данные!F4": 14400, "Данные!G4": 57600, "Summary!B3": 333630}},
        "chart": True, "chart_value_column": "G", "text": ["Скидка", "После скидки, ₽"],
    })

    aggregates = copy.deepcopy(base)
    aggregates["chart"] = None
    aggregates["summary"] = [
        {"label": "Итого, ₽", "operation": "sum", "column": "cost"},
        {"label": "Средняя стоимость, ₽", "operation": "average", "column": "cost"},
        {"label": "Минимум, ₽", "operation": "min", "column": "cost"},
        {"label": "Максимум, ₽", "operation": "max", "column": "cost"},
        {"label": "Количество услуг", "operation": "count", "column": "service"},
    ]
    cases["aggregate_stats"] = (aggregates, {
        "sheets": ["Summary", "Данные"],
        "formulas": {"Summary!B3": "=SUM('Данные'!D4:D9)", "Summary!B7": "=COUNTA('Данные'!A4:A9)"},
        "values": {"Summary!B3": 378700, "Summary!B4": 378700 / 6,
                   "Summary!B5": 34800, "Summary!B6": 86400, "Summary!B7": 6},
        "chart": False, "text": ["Средняя стоимость, ₽", "Максимум, ₽"],
    })

    trend = copy.deepcopy(base)
    trend["table"]["columns"][0]["header"] = "Месяц"
    trend["table"]["columns"].append({"key": "extra", "header": "Дополнительно, ₽", "type": "number"})
    trend["table"]["calculatedColumns"].append({
        "key": "total", "header": "Всего, ₽", "operation": "add", "inputs": ["cost", "extra"],
    })
    for row, month in zip(trend["table"]["rows"],
                          ("Январь", "Февраль", "Март", "Апрель", "Май", "Июнь")):
        row["service"] = month
        row["extra"] = 1000
    trend["chart"].update(kind="line", title="Стоимость по месяцам, ₽", value="total")
    trend["summary"][0]["column"] = "total"
    cases["line_trend"] = (trend, {
        "sheets": ["Summary", "Данные"],
        "formulas": {"Данные!E4": "=B4*C4", "Данные!F4": "=E4+D4",
                     "Summary!B3": "=SUM('Данные'!F4:F9)"},
        "values": {"Данные!F4": 73000, "Summary!B3": 384700},
        "chart": True, "chart_kind": "line", "text": ["Месяц", "Стоимость по месяцам, ₽"],
    })

    chained = copy.deepcopy(base)
    chained["chart"] = None
    chained["table"]["columns"][2]["header"] = "Цена, ₽"
    chained["table"]["columns"].append({
        "key": "discount", "header": "Скидка", "type": "percent", "minimum": 0, "maximum": 1,
    })
    for row in chained["table"]["rows"]:
        row["discount"] = 0.1
    chained["table"]["calculatedColumns"].extend([
        {"key": "discount_amount", "header": "Скидка, ₽",
         "operation": "multiply", "inputs": ["cost", "discount"]},
        {"key": "net_cost", "header": "К оплате, ₽",
         "operation": "subtract", "inputs": ["cost", "discount_amount"]},
        {"key": "net_unit", "header": "За ед., ₽",
         "operation": "divide", "inputs": ["net_cost", "units"]},
    ])
    chained["summary"][0]["column"] = "net_cost"
    cases["chained_calculations"] = (chained, {
        "sheets": ["Summary", "Данные"],
        "formulas": {"Данные!E4": "=B4*C4", "Данные!F4": "=E4*D4",
                     "Данные!G4": "=E4-F4", "Данные!H4": "=G4/B4"},
        "values": {"Данные!F4": 7200, "Данные!G4": 64800,
                   "Данные!H4": 16200, "Summary!B3": 340830},
        "chart": False, "text": ["К оплате, ₽", "За ед., ₽"],
    })

    literal = copy.deepcopy(base)
    literal["chart"] = None
    literal["table"]["rows"][0]["service"] = "=1+1"
    cases["literal_source_text"] = (literal, {
        "sheets": ["Summary", "Данные"],
        "formulas": {"Данные!D4": "=B4*C4"},
        "values": {"Данные!D4": 72000, "Summary!B3": 378700},
        "literal": ("Данные!A4", "=1+1"),
        "chart": False, "text": ["=1+1", "Стоимость, ₽"],
    })

    long_table = copy.deepcopy(base)
    long_table["summary"] = []
    long_table["chart"] = None
    long_table["table"]["rows"] = [
        {"service": f"Service {index:02d}", "units": index, "unit_cost": 125}
        for index in range(1, 81)
    ]
    cases["long_table"] = (long_table, {
        "sheets": ["Данные"],
        "formulas": {"Данные!D4": "=B4*C4", "Данные!D83": "=B83*C83"},
        "values": {"Данные!D4": 125, "Данные!D83": 10000},
        "chart": False, "text": ["Service 01", "Service 80"],
        "pages_min": 2, "pages_max": 5, "header_every_page": "Цена за единицу, ₽",
    })
    return cases


def _sheet_cell(workbook, address: str):
    sheet, cell = address.split("!", 1)
    return workbook[sheet][cell]


def _require(condition: bool, message: object) -> None:
    if not condition:
        raise AssertionError(message)


def _render_pixels(pdf: Path, folder: Path) -> tuple[str, ...]:
    result = subprocess.run(
        ["pdftoppm", "-r", "72", "-png", str(pdf), str(folder / "page")],
        capture_output=True, text=True, timeout=120, check=False,
    )
    if result.returncode:
        raise AssertionError(f"PDF rasterization failed: {result.stderr}")
    fingerprints = []
    for path in sorted(folder.glob("page-*.png")):
        with Image.open(path) as image:
            fingerprints.append(hashlib.sha256(image.convert("RGB").tobytes()).hexdigest())
    if not fingerprints:
        raise AssertionError("No PDF pages rendered")
    return tuple(fingerprints)


def _verify_mutation(output: Path, oracle: dict, run_dir: Path) -> None:
    mutation = oracle.get("mutation")
    if not mutation:
        return
    original_hash = hashlib.sha256(output.read_bytes()).digest()
    edited = run_dir / "edited-copy.xlsx"
    workbook = load_workbook(output)
    try:
        address, replacement = mutation["input"]
        _sheet_cell(workbook, address).value = replacement
        if oracle["chart"]:
            chart = workbook["Summary"]._charts[0]
            value_reference = chart.series[0].val.numRef.f
            column = oracle["chart_value_column"]
            _require(value_reference == f"'Данные'!${column}$4:${column}$9",
                     ("Chart is not bound to editable source cells", value_reference))
        workbook.save(edited)
    finally:
        workbook.close()
    recalculated_dir = run_dir / "edited-recalc"
    recalculated_dir.mkdir()
    result = subprocess.run(
        ["soffice", f"-env:UserInstallation={(run_dir / 'edited-profile').as_uri()}",
         "--headless", "--convert-to", "xlsx", "--outdir", str(recalculated_dir), str(edited)],
        capture_output=True, text=True, timeout=120, check=False,
    )
    recalculated = recalculated_dir / edited.name
    _require(result.returncode == 0 and recalculated.is_file(),
             ("Edited workbook did not recalculate", result.stderr or result.stdout))
    values = load_workbook(recalculated, data_only=True)
    try:
        for address, expected in mutation["values"].items():
            actual = _sheet_cell(values, address).value
            _require(isinstance(actual, (int, float)) and not isinstance(actual, bool) and
                     math.isclose(actual, expected, rel_tol=1e-9), (address, actual, expected))
        for sheet in values:
            for row in sheet:
                for cell in row:
                    _require(cell.data_type != "e", (sheet.title, cell.coordinate, cell.value))
    finally:
        values.close()
    _require(hashlib.sha256(output.read_bytes()).digest() == original_hash,
             "Testing an input edit changed the delivered workbook")


def _verify(case_id: str, spec: dict, oracle: dict, run_dir: Path) -> tuple[str, ...]:
    run_dir.mkdir(parents=True)
    spec["job"]["filename"] = f"{case_id}.xlsx"
    spec_path = run_dir / "spec.json"
    spec_path.write_text(json.dumps(spec, ensure_ascii=False), encoding="utf-8")
    output = run_dir / spec["job"]["filename"]
    result = subprocess.run(
        [sys.executable, str(BUILDER), str(spec_path), str(output)],
        capture_output=True, text=True, timeout=180, check=False,
    )
    if result.returncode:
        raise AssertionError(f"Builder failed: {result.stderr}")
    report = json.loads(Path(f"{output}.artifact-report.json").read_text(encoding="utf-8"))
    if report["status"] != "ready" or report["issues"] or any(
        check["status"] != "passed" for check in report["qaChecks"]
    ):
        raise AssertionError(f"QA report is not green: {report}")
    workbook = load_workbook(output, data_only=False)
    try:
        _require(workbook.sheetnames == oracle["sheets"], workbook.sheetnames)
        data = workbook[spec["table"]["sheetName"]]
        _require(spec["table"]["name"] in data.tables, "Native Excel Table is missing")
        _require(data.freeze_panes == "A4", "Freeze pane is missing")
        chart_present = bool(workbook["Summary"]._charts) if "Summary" in workbook else False
        _require(chart_present == oracle["chart"], "Native chart presence differs")
        if chart_present:
            expected_chart = LineChart if oracle.get("chart_kind") == "line" else BarChart
            _require(isinstance(workbook["Summary"]._charts[0], expected_chart),
                     "Native chart type differs")
        for address, formula in oracle["formulas"].items():
            _require(_sheet_cell(workbook, address).value == formula, address)
        if "date" in oracle:
            address, expected_date = oracle["date"]
            _require(_sheet_cell(workbook, address).value == expected_date, address)
        if "literal" in oracle:
            address, expected_text = oracle["literal"]
            cell = _sheet_cell(workbook, address)
            _require(cell.value == expected_text and cell.data_type == "s", address)
    finally:
        workbook.close()

    recalculated_dir = run_dir / "oracle-recalc"
    recalculated_dir.mkdir()
    result = subprocess.run(
        [
            "soffice", f"-env:UserInstallation={(run_dir / 'oracle-profile').as_uri()}",
            "--headless", "--convert-to", "xlsx", "--outdir", str(recalculated_dir), str(output),
        ],
        capture_output=True, text=True, timeout=120, check=False,
    )
    recalculated_path = recalculated_dir / output.name
    if result.returncode or not recalculated_path.is_file():
        raise AssertionError(f"Independent recalculation failed: {result.stderr or result.stdout}")
    values = load_workbook(recalculated_path, data_only=True)
    try:
        for address, expected in oracle["values"].items():
            actual = _sheet_cell(values, address).value
            _require(
                isinstance(actual, (int, float)) and math.isclose(actual, expected, rel_tol=1e-9),
                (address, actual, expected),
            )
        for sheet in values:
            for row in sheet:
                for cell in row:
                    _require(cell.data_type != "e", (sheet.title, cell.coordinate, cell.value))
    finally:
        values.close()

    _verify_mutation(output, oracle, run_dir)

    pdf = run_dir / f"_qa_{output.stem}-preview.pdf"
    rendered_text = subprocess.run(
        ["pdftotext", "-layout", str(pdf), "-"],
        capture_output=True, text=True, timeout=30, check=True,
    ).stdout
    for phrase in oracle["text"]:
        _require(phrase in rendered_text, (case_id, phrase))
    pages = _render_pixels(pdf, run_dir)
    _require(oracle.get("pages_min", len(oracle["sheets"])) <= len(pages) <=
             oracle.get("pages_max", len(oracle["sheets"])),
             (case_id, len(pages), len(oracle["sheets"])))
    if "header_every_page" in oracle:
        for page in range(1, len(pages) + 1):
            page_text = subprocess.run(
                ["pdftotext", "-f", str(page), "-l", str(page), "-layout", str(pdf), "-"],
                capture_output=True, text=True, timeout=30, check=True,
            ).stdout
            _require(oracle["header_every_page"] in page_text, (case_id, page, "missing header"))
    return pages


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--runs", type=int, default=1)
    parser.add_argument("--output-dir", type=Path)
    args = parser.parse_args()
    if not 1 <= args.runs <= 3:
        parser.error("--runs must be between 1 and 3")
    if args.output_dir and args.output_dir.exists():
        parser.error("--output-dir must not already exist")
    temporary = tempfile.TemporaryDirectory(prefix="xlsx-core-goldens-") if not args.output_dir else None
    try:
        output_dir = args.output_dir or Path(temporary.name)
        output_dir.mkdir(parents=True, exist_ok=True)
        for case_id, (spec, oracle) in _cases().items():
            fingerprints = []
            for run in range(1, args.runs + 1):
                fingerprints.append(_verify(case_id, copy.deepcopy(spec), oracle, output_dir / case_id / f"run-{run}"))
            if len(set(fingerprints)) != 1:
                raise AssertionError(f"Rendered pixels changed between {case_id} runs")
            print(f"{case_id}: {args.runs}/{args.runs} passed")
        print(
            f"Core-only evaluation passed: {len(_cases())} cases × {args.runs} runs; "
            "not the mixed-mode XLSX acceptance gate"
        )
        if args.output_dir:
            print(f"Rendered evidence: {output_dir}")
        return 0
    finally:
        if temporary:
            temporary.cleanup()


if __name__ == "__main__":
    raise SystemExit(main())
