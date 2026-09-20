"""Structural and real-office checks for the focused XLSX builder."""

from __future__ import annotations

import copy
import json
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from skill.xlsx.scripts.build_spreadsheet import SpecError, _office_convert, _validate


ROOT = Path(__file__).resolve().parents[2]
BUILDER = ROOT / "skill/xlsx/scripts/build_spreadsheet.py"
RUSSIAN_FIXTURE = ROOT / "bench/spreadsheets/fixtures/ru_service_cost.json"


def example_spec() -> dict:
    return {
        "job": {
            "format": "xlsx",
            "audience": "Operations lead",
            "goal": "Review service cost",
            "sourceFileIds": ["pilot-notes.csv"],
            "immutableElements": [],
            "locale": "ru-RU",
            "filename": "service-cost.xlsx",
            "acceptanceCriteria": ["Formulas update", "Every sheet renders"],
        },
        "title": "Service cost",
        "table": {
            "sheetName": "Data",
            "name": "ServiceCosts",
            "columns": [
                {"key": "service", "header": "Service", "type": "text", "choices": ["Support", "Training"]},
                {"key": "units", "header": "Units", "type": "integer", "minimum": 0},
                {"key": "unit_cost", "header": "Unit cost", "type": "number", "minimum": 0},
            ],
            "rows": [
                {"service": "Support", "units": 12, "unit_cost": 125},
                {"service": "Training", "units": 4, "unit_cost": 300},
            ],
            "calculatedColumns": [
                {"key": "cost", "header": "Cost", "operation": "multiply", "inputs": ["units", "unit_cost"]}
            ],
        },
        "summary": [{"label": "Total cost", "operation": "sum", "column": "cost"}],
        "chart": {"kind": "bar", "title": "Cost by service", "category": "service", "value": "cost"},
        "sources": [{"label": "Pilot data", "location": "pilot-notes.csv"}],
        "changeLog": [],
        "repairIterations": 0,
    }


class SpreadsheetBuilderTests(unittest.TestCase):
    def test_russian_fixture_renders_localized_summary_and_chart(self):
        with tempfile.TemporaryDirectory() as folder:
            output = Path(folder) / "stoimost-uslug.xlsx"
            result = subprocess.run(
                [sys.executable, str(BUILDER), str(RUSSIAN_FIXTURE), str(output)],
                capture_output=True, text=True, timeout=180, check=False,
            )
            self.assertEqual(result.returncode, 0, result.stderr)
            report = json.loads(Path(f"{output}.artifact-report.json").read_text(encoding="utf-8"))
            self.assertEqual(report["status"], "ready", report)
            pdf = output.with_name("_qa_stoimost-uslug-preview.pdf")
            rendered = subprocess.run(
                ["pdftotext", "-layout", str(pdf), "-"],
                capture_output=True, text=True, timeout=30, check=True,
            ).stdout
            self.assertIn("Показатель", rendered)
            self.assertIn("Значение", rendered)
            self.assertIn("Стоимость по услугам", rendered)
            self.assertIn("Источник: Демонстрационные данные", rendered)
            self.assertNotIn("Metric", rendered)

    def test_real_office_roundtrip_keeps_formula_table_chart_and_values(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            source = root / "pilot-notes.csv"
            source.write_bytes(b"original source bytes")
            spec_path = root / "spec.json"
            spec_path.write_text(json.dumps(example_spec()), encoding="utf-8")
            output = root / "service-cost.xlsx"
            result = subprocess.run(
                [sys.executable, str(BUILDER), str(spec_path), str(output)],
                capture_output=True,
                text=True,
                timeout=180,
                check=False,
            )
            self.assertEqual(result.returncode, 0, result.stderr)
            report = json.loads(Path(f"{output}.artifact-report.json").read_text(encoding="utf-8"))
            self.assertEqual(report["status"], "ready", report)
            self.assertEqual(report["sourceFileIds"], ["pilot-notes.csv"])
            self.assertEqual(source.read_bytes(), b"original source bytes")
            self.assertTrue((root / "_qa_service-cost-preview.pdf").is_file())
            workbook = load_workbook(output, data_only=False)
            self.assertEqual(workbook.sheetnames, ["Summary", "Data"])
            self.assertEqual(workbook["Data"]["D4"].value, "=B4*C4")
            self.assertEqual(workbook["Summary"]["B3"].value, "=SUM('Data'!D4:D5)")
            self.assertIn("ServiceCosts", workbook["Data"].tables)
            self.assertEqual(len(workbook["Summary"]._charts), 1)
            self.assertTrue(workbook["Data"].data_validations.dataValidation)
            workbook["Data"]["B4"] = 20
            workbook.save(output)
            workbook.close()
            recalculated = _office_convert(output, root / "updated", "xlsx", root / "updated-profile")
            values = load_workbook(recalculated, data_only=True)
            self.assertEqual(values["Data"]["D4"].value, 2500)
            self.assertEqual(values["Summary"]["B3"].value, 3700)
            values.close()

    def test_source_text_that_looks_like_a_formula_remains_literal(self):
        spec = example_spec()
        spec["table"]["columns"][0].pop("choices")
        spec["table"]["rows"][0]["service"] = '=WEBSERVICE("https://example.invalid")'
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            spec_path = root / "spec.json"
            spec_path.write_text(json.dumps(spec), encoding="utf-8")
            output = root / "service-cost.xlsx"
            result = subprocess.run(
                [sys.executable, str(BUILDER), str(spec_path), str(output)],
                capture_output=True,
                text=True,
                timeout=180,
                check=False,
            )
            self.assertEqual(result.returncode, 0, result.stderr)
            workbook = load_workbook(output, data_only=False)
            self.assertEqual(workbook["Data"]["A4"].data_type, "s")
            self.assertEqual(workbook["Data"]["A4"].value, '=WEBSERVICE("https://example.invalid")')
            workbook.close()

    def test_zero_divisor_rejected_before_file_creation(self):
        spec = example_spec()
        spec["table"]["calculatedColumns"][0]["operation"] = "divide"
        spec["table"]["rows"][1]["unit_cost"] = 0
        with self.assertRaisesRegex(SpecError, "zero divisor"):
            _validate(spec, Path("service-cost.xlsx"))

    def test_raw_formulas_and_template_paths_are_not_accepted(self):
        for field, value in (("rawFormulas", ["=WEBSERVICE(...)"],), ("templatePath", "/tmp/untrusted.xlsx")):
            with self.subTest(field=field):
                spec = example_spec()
                spec[field] = value
                with self.assertRaisesRegex(SpecError, "not supported"):
                    _validate(spec, Path("service-cost.xlsx"))

    def test_missing_numeric_input_is_not_silently_zero(self):
        spec = example_spec()
        spec["table"]["rows"][0]["units"] = None
        with self.assertRaisesRegex(SpecError, "finite number"):
            _validate(spec, Path("service-cost.xlsx"))

    def test_wide_worksheet_is_rejected_before_unreadable_render(self):
        spec = example_spec()
        for index in range(5):
            key = f"extra_{index}"
            spec["table"]["columns"].append({"key": key, "header": key, "type": "integer"})
            for row in spec["table"]["rows"]:
                row[key] = 1
        with self.assertRaisesRegex(SpecError, "1 to 8 total columns"):
            _validate(spec, Path("service-cost.xlsx"))

    def test_original_spec_is_not_modified_by_validation(self):
        spec = example_spec()
        original = copy.deepcopy(spec)
        _validate(spec, Path("service-cost.xlsx"))
        self.assertEqual(spec, original)

    def test_text_count_uses_counta_not_numeric_count(self):
        spec = example_spec()
        spec["summary"] = [{"label": "Number of services", "operation": "count", "column": "service"}]
        spec["chart"] = None
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            spec_path = root / "spec.json"
            spec_path.write_text(json.dumps(spec), encoding="utf-8")
            output = root / "service-cost.xlsx"
            result = subprocess.run(
                [sys.executable, str(BUILDER), str(spec_path), str(output)],
                capture_output=True, text=True, timeout=180, check=False,
            )
            self.assertEqual(result.returncode, 0, result.stderr)
            report = json.loads(Path(f"{output}.artifact-report.json").read_text(encoding="utf-8"))
            self.assertEqual(report["status"], "ready", report)
            workbook = load_workbook(output, data_only=False)
            self.assertEqual(workbook["Summary"]["B3"].value, "=COUNTA('Data'!A4:A5)")
            workbook.close()

    def test_invalid_excel_identifiers_and_lists_are_rejected(self):
        cases = [
            (lambda spec: spec["table"].update(name="A1"), "cell coordinate"),
            (lambda spec: spec["table"].update(name="R"), "reserved name"),
            (lambda spec: spec["table"].update(name="R1C1"), "Excel reference"),
            (lambda spec: spec["table"].update(sheetName="Data!More"), "valid Excel sheet name"),
            (lambda spec: spec["table"]["columns"][0].update(choices=["A,B"]), "cannot contain"),
            (lambda spec: spec["table"]["calculatedColumns"][0].update(numberFormat=";;;"), "numberFormat is unsupported"),
            (lambda spec: spec["table"]["columns"][0].update(type=[]), "unsupported type"),
            (lambda spec: spec["table"]["columns"][0].update(choices=[{}]), "invalid or duplicate"),
            (lambda spec: spec["table"]["calculatedColumns"][0].update(operation=[]), "invalid header or operation"),
            (lambda spec: spec["table"]["calculatedColumns"][0].update(inputs=[{}, "units"]), "preceding numeric keys"),
            (lambda spec: spec["summary"][0].update(operation=[]), "unsupported operation"),
            (lambda spec: spec.update(chart={"kind": [], "title": "Test"}), "chart.kind"),
            (lambda spec: spec.update(repairIterations=True), "repairIterations"),
        ]
        for change, error in cases:
            with self.subTest(error=error):
                spec = example_spec()
                change(spec)
                with self.assertRaisesRegex(SpecError, error):
                    _validate(spec, Path("service-cost.xlsx"))

    def test_existing_output_is_never_overwritten_even_during_repair(self):
        for repair_iteration in (0, 1):
            with self.subTest(repair_iteration=repair_iteration), tempfile.TemporaryDirectory() as folder:
                spec = example_spec()
                spec["repairIterations"] = repair_iteration
                root = Path(folder)
                spec_path = root / "spec.json"
                spec_path.write_text(json.dumps(spec), encoding="utf-8")
                output = root / "service-cost.xlsx"
                output.write_bytes(b"existing workbook bytes")
                result = subprocess.run(
                    [sys.executable, str(BUILDER), str(spec_path), str(output)],
                    capture_output=True, text=True, timeout=30, check=False,
                )
                self.assertNotEqual(result.returncode, 0)
                self.assertIn("cannot be overwritten", result.stderr)
                self.assertEqual(output.read_bytes(), b"existing workbook bytes")


if __name__ == "__main__":
    unittest.main()
