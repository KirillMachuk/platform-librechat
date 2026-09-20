#!/usr/bin/env python3
"""Build a focused editable XLSX and verify formulas in an office renderer."""

from __future__ import annotations

import json
import math
import re
import shutil
import subprocess
import sys
import tempfile
import urllib.parse
from datetime import date
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.utils.cell import coordinate_from_string
from openpyxl.utils.exceptions import CellCoordinatesException
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.properties import PageSetupProperties
from openpyxl.worksheet.page import PageMargins
from openpyxl.worksheet.table import Table, TableStyleInfo
from PIL import Image, ImageChops


SKILL_VERSION = "0.1.0"
MAX_ROWS = 200
MAX_COLUMNS = 8
MAX_REPAIR_ITERATIONS = 2
FONT_NAME = "Liberation Sans"
BASE_TYPES = {"text", "integer", "number", "percent", "date"}
CALCULATIONS = {"add", "subtract", "multiply", "divide"}
SUMMARIES = {"sum", "average", "min", "max", "count"}
COLORS = {
    "ink": "20242A",
    "muted": "606A78",
    "accent": "2458A6",
    "surface": "F3F6FA",
    "line": "D8DEE8",
    "warning": "FCE8E6",
}


class SpecError(ValueError):
    """The requested workbook cannot be built without guessing."""


def _text(value: Any, field: str, limit: int = 500) -> str:
    if not isinstance(value, str) or not value.strip() or len(value) > limit:
        raise SpecError(f"{field} must be non-empty text of at most {limit} characters")
    if any(ord(char) < 32 and char not in "\t\n" for char in value):
        raise SpecError(f"{field} contains a control character")
    return value.strip()


def _items(value: Any, field: str, limit: int) -> list[Any]:
    if not isinstance(value, list) or len(value) > limit:
        raise SpecError(f"{field} must be a list with at most {limit} items")
    return value


def _number(value: Any, field: str, integer: bool = False) -> int | float:
    if isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value):
        raise SpecError(f"{field} must be a finite number")
    if integer and (not isinstance(value, int) or isinstance(value, bool)):
        raise SpecError(f"{field} must be an integer")
    if abs(value) > 10**15:
        raise SpecError(f"{field} exceeds the supported numeric range")
    return value


def _sheet_name(value: Any, field: str) -> str:
    name = _text(value, field, 31)
    if re.search(r"[\\/*?:!\[\]]", name) or name.startswith("'") or name.endswith("'"):
        raise SpecError(f"{field} is not a valid Excel sheet name")
    if name.casefold() in {"summary", "sources"}:
        raise SpecError(f"{field} conflicts with a reserved output sheet")
    return name


def _source(source: Any, index: int) -> dict[str, str]:
    if not isinstance(source, dict):
        raise SpecError(f"sources[{index}] must be an object")
    label = _text(source.get("label"), f"sources[{index}].label", 160)
    location = source.get("location")
    url = source.get("url")
    if bool(location) == bool(url):
        raise SpecError(f"sources[{index}] needs exactly one location or url")
    if location:
        return {"label": label, "location": _text(location, f"sources[{index}].location", 300)}
    url = _text(url, f"sources[{index}].url", 500)
    parsed = urllib.parse.urlsplit(url)
    if parsed.scheme not in {"http", "https"} or not parsed.hostname or parsed.username or parsed.password:
        raise SpecError(f"sources[{index}].url must be an absolute credential-free web URL")
    if any(char.isspace() for char in url) or "\\" in url:
        raise SpecError(f"sources[{index}].url contains whitespace or a backslash")
    return {"label": label, "url": url}


def _typed_value(value: Any, column: dict[str, Any], row_index: int) -> Any:
    field = f"table.rows[{row_index}].{column['key']}"
    kind = column["type"]
    if kind == "text":
        result = _text(value, field)
        choices = column.get("choices")
        if choices and result not in choices:
            raise SpecError(f"{field} is not among its declared choices")
        return result
    if kind == "date":
        raw = _text(value, field, 10)
        try:
            return date.fromisoformat(raw)
        except ValueError as exc:
            raise SpecError(f"{field} must be an ISO date") from exc
    result = _number(value, field, kind == "integer")
    if kind == "percent" and not 0 <= result <= 1:
        raise SpecError(f"{field} must be a fraction between 0 and 1")
    if "minimum" in column and result < column["minimum"]:
        raise SpecError(f"{field} is below its declared minimum")
    if "maximum" in column and result > column["maximum"]:
        raise SpecError(f"{field} is above its declared maximum")
    return result


def _calculate(operation: str, left: int | float, right: int | float, field: str) -> int | float:
    if operation == "divide" and right == 0:
        raise SpecError(f"{field} has a zero divisor")
    result = {
        "add": lambda: left + right,
        "subtract": lambda: left - right,
        "multiply": lambda: left * right,
        "divide": lambda: left / right,
    }[operation]()
    return _number(result, field)


def _validate(spec: Any, output: Path) -> dict[str, Any]:
    if not isinstance(spec, dict):
        raise SpecError("The specification must be an object")
    for unsupported in ("inputPath", "templatePath", "rawFormulas", "outputPdf"):
        if unsupported in spec:
            raise SpecError(f"{unsupported} is not supported by this builder version")
    job = spec.get("job")
    if not isinstance(job, dict) or job.get("format") != "xlsx":
        raise SpecError("job.format must be xlsx")
    for field in ("audience", "goal", "locale", "filename"):
        _text(job.get(field), f"job.{field}", 180)
    if job["filename"] != output.name or output.suffix.lower() != ".xlsx":
        raise SpecError("Output path must match job.filename and end in .xlsx")
    for field in ("sourceFileIds", "immutableElements", "acceptanceCriteria"):
        values = _items(job.get(field), f"job.{field}", 100)
        if field == "acceptanceCriteria" and not values:
            raise SpecError("job.acceptanceCriteria cannot be empty")
        for index, value in enumerate(values):
            _text(value, f"job.{field}[{index}]", 300)
    title = _text(spec.get("title"), "title", 160)
    table = spec.get("table")
    if not isinstance(table, dict):
        raise SpecError("table must be an object")
    sheet_name = _sheet_name(table.get("sheetName"), "table.sheetName")
    table_name = _text(table.get("name"), "table.name", 80)
    if not re.fullmatch(r"[A-Za-z_][A-Za-z0-9_]*", table_name):
        raise SpecError("table.name must be an Excel-safe identifier")
    if table_name.casefold() in {"r", "c"} or re.fullmatch(
        r"R[1-9][0-9]*C[1-9][0-9]*", table_name, re.IGNORECASE
    ):
        raise SpecError("table.name cannot be an Excel reference or reserved name")
    try:
        coordinate_from_string(table_name)
    except CellCoordinatesException:
        pass
    else:
        raise SpecError("table.name cannot be a cell coordinate")
    columns = _items(table.get("columns"), "table.columns", MAX_COLUMNS)
    calculated = _items(table.get("calculatedColumns", []), "table.calculatedColumns", MAX_COLUMNS)
    if not columns or len(columns) + len(calculated) > MAX_COLUMNS:
        raise SpecError("The table needs 1 to 8 total columns in this builder version")
    types: dict[str, str] = {}
    headers: set[str] = set()
    for index, column in enumerate(columns):
        if not isinstance(column, dict):
            raise SpecError(f"table.columns[{index}] must be an object")
        key = _text(column.get("key"), f"table.columns[{index}].key", 40)
        header = _text(column.get("header"), f"table.columns[{index}].header", 100)
        kind = column.get("type")
        if not re.fullmatch(r"[A-Za-z_][A-Za-z0-9_]*", key) or key in types:
            raise SpecError(f"table.columns[{index}].key must be a unique identifier")
        if header.casefold() in headers or not isinstance(kind, str) or kind not in BASE_TYPES:
            raise SpecError(f"table.columns[{index}] has a duplicate header or unsupported type")
        types[key] = kind
        headers.add(header.casefold())
        choices = column.get("choices")
        if choices is not None:
            if kind != "text" or not isinstance(choices, list) or not 1 <= len(choices) <= 20:
                raise SpecError(f"table.columns[{index}].choices needs a text column and 1 to 20 values")
            if any(not isinstance(choice, str) for choice in choices) or len(set(choices)) != len(choices):
                raise SpecError(f"table.columns[{index}].choices contains invalid or duplicate values")
            for choice_index, choice in enumerate(choices):
                _text(choice, f"table.columns[{index}].choices[{choice_index}]", 50)
                if any(char in choice for char in ',"\n\r'):
                    raise SpecError(f"table.columns[{index}].choices cannot contain comma, quote, or newline")
            if len(",".join(choices)) > 250:
                raise SpecError(f"table.columns[{index}].choices exceeds the Excel list limit")
        for bound in ("minimum", "maximum"):
            if bound in column:
                if kind not in {"integer", "number", "percent"}:
                    raise SpecError(f"{bound} requires a numeric column")
                _number(column[bound], f"table.columns[{index}].{bound}")
        if "minimum" in column and "maximum" in column and column["minimum"] > column["maximum"]:
            raise SpecError(f"table.columns[{index}] has reversed numeric bounds")
    for index, column in enumerate(calculated):
        if not isinstance(column, dict):
            raise SpecError(f"table.calculatedColumns[{index}] must be an object")
        key = _text(column.get("key"), f"table.calculatedColumns[{index}].key", 40)
        header = _text(column.get("header"), f"table.calculatedColumns[{index}].header", 100)
        inputs = column.get("inputs")
        if not re.fullmatch(r"[A-Za-z_][A-Za-z0-9_]*", key) or key in types:
            raise SpecError(f"table.calculatedColumns[{index}].key must be unique")
        operation = column.get("operation")
        if header.casefold() in headers or not isinstance(operation, str) or operation not in CALCULATIONS:
            raise SpecError(f"table.calculatedColumns[{index}] has an invalid header or operation")
        if not isinstance(inputs, list) or len(inputs) != 2 or any(
            not isinstance(operand, str)
            or operand not in types
            or types[operand] not in {"integer", "number", "percent", "calculated"}
            for operand in inputs
        ):
            raise SpecError(f"table.calculatedColumns[{index}].inputs must name two preceding numeric keys")
        if "numberFormat" in column:
            if not isinstance(column["numberFormat"], str) or column["numberFormat"] not in {
                "#,##0", "#,##0.00", "0.0%", "0.00%"
            }:
                raise SpecError(f"table.calculatedColumns[{index}].numberFormat is unsupported")
        if "highlightNegative" in column and not isinstance(column["highlightNegative"], bool):
            raise SpecError(f"table.calculatedColumns[{index}].highlightNegative must be boolean")
        types[key] = "calculated"
        headers.add(header.casefold())
    rows = _items(table.get("rows"), "table.rows", MAX_ROWS)
    if not rows:
        raise SpecError("table.rows cannot be empty")
    expected_rows: list[dict[str, Any]] = []
    base_keys = {column["key"] for column in columns}
    for row_index, row in enumerate(rows):
        if not isinstance(row, dict) or set(row) != base_keys:
            raise SpecError(f"table.rows[{row_index}] must contain exactly the base column keys")
        values = {column["key"]: _typed_value(row[column["key"]], column, row_index) for column in columns}
        for column in calculated:
            left, right = (values[key] for key in column["inputs"])
            values[column["key"]] = _calculate(column["operation"], left, right, f"table.rows[{row_index}].{column['key']}")
        expected_rows.append(values)
    summary = _items(spec.get("summary", []), "summary", 12)
    for index, item in enumerate(summary):
        if not isinstance(item, dict):
            raise SpecError(f"summary[{index}] must be an object")
        _text(item.get("label"), f"summary[{index}].label", 100)
        if (
            not isinstance(item.get("operation"), str)
            or item["operation"] not in SUMMARIES
            or not isinstance(item.get("column"), str)
            or item["column"] not in types
        ):
            raise SpecError(f"summary[{index}] has an unsupported operation or column")
        if item["operation"] != "count" and types[item["column"]] in {"text", "date"}:
            raise SpecError(f"summary[{index}] needs a numeric column")
    chart = spec.get("chart")
    if chart is not None:
        if not isinstance(chart, dict) or not isinstance(chart.get("kind"), str) or chart["kind"] not in {
            "bar", "line"
        }:
            raise SpecError("chart.kind must be bar or line")
        _text(chart.get("title"), "chart.title", 100)
        if not isinstance(chart.get("category"), str) or not isinstance(chart.get("value"), str):
            raise SpecError("chart needs a text category and numeric value column")
        if types.get(chart["category"]) != "text" or types.get(chart["value"]) not in {
            "integer", "number", "percent", "calculated"
        }:
            raise SpecError("chart needs a text category and numeric value column")
    sources = [_source(item, index) for index, item in enumerate(_items(spec.get("sources", []), "sources", 20))]
    changes = _items(spec.get("changeLog", []), "changeLog", 20)
    for index, change in enumerate(changes):
        if not isinstance(change, dict):
            raise SpecError(f"changeLog[{index}] must be an object")
        _text(change.get("target"), f"changeLog[{index}].target", 100)
        _text(change.get("summary"), f"changeLog[{index}].summary", 300)
    repair_iterations = spec.get("repairIterations", 0)
    if isinstance(repair_iterations, bool) or not isinstance(repair_iterations, int) or not 0 <= repair_iterations <= MAX_REPAIR_ITERATIONS:
        raise SpecError("repairIterations must be between 0 and 2")
    return {
        "job": job,
        "title": title,
        "table": table,
        "sheetName": sheet_name,
        "columns": columns,
        "calculated": calculated,
        "rows": expected_rows,
        "summary": summary,
        "chart": chart,
        "sources": sources,
        "changes": changes,
        "repairIterations": repair_iterations,
    }


def _literal(cell: Any, value: str) -> None:
    cell.value = value
    cell.data_type = "s"


def _number_format(kind: str) -> str:
    return {"integer": "#,##0", "number": "#,##0.00", "percent": "0.0%", "date": "yyyy-mm-dd"}.get(kind, "General")


def _style_header(sheet: Any, row: int, width: int) -> None:
    for cells in sheet.iter_rows(min_row=row, max_row=row, min_col=1, max_col=width):
        for cell in cells:
            cell.fill = PatternFill("solid", fgColor=COLORS["accent"])
            cell.font = Font(name=FONT_NAME, size=10, bold=True, color="FFFFFF")
            cell.alignment = Alignment(horizontal="center", vertical="center")
    sheet.row_dimensions[row].height = 26


def _build(spec: dict[str, Any], output: Path) -> tuple[dict[str, str], dict[str, int | float], list[str]]:
    russian = spec["job"]["locale"].lower().startswith("ru")
    has_summary = bool(spec["summary"] or spec["chart"])
    workbook = Workbook()
    if has_summary:
        summary_sheet = workbook.active
        summary_sheet.title = "Summary"
        data = workbook.create_sheet(spec["sheetName"])
    else:
        summary_sheet = None
        data = workbook.active
        data.title = spec["sheetName"]
    data.sheet_view.showGridLines = False
    data.freeze_panes = "A4"
    data.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    data.page_setup.fitToWidth = 1
    data.page_setup.fitToHeight = 0
    data.page_setup.orientation = "landscape"
    data.page_setup.paperSize = data.PAPERSIZE_A4
    data.page_margins = PageMargins(left=0.4, right=0.4, top=0.5, bottom=0.5, header=0.2, footer=0.2)
    data.print_options.horizontalCentered = True
    data.sheet_properties.outlinePr.summaryBelow = True
    _literal(data["A1"], spec["title"])
    data["A1"].font = Font(name=FONT_NAME, size=14, bold=True, color=COLORS["ink"])
    if len(spec["sources"]) == 1:
        source = spec["sources"][0]
        _literal(data["A2"], f"{'Источник' if russian else 'Source'}: {source['label']} — {source.get('location') or source.get('url')}")
    elif spec["sources"]:
        _literal(data["A2"], "Источники: см. лист Sources" if russian else "Sources: see the Sources sheet")
    data["A2"].font = Font(name=FONT_NAME, size=9, italic=True, color=COLORS["muted"])
    all_columns = [*spec["columns"], *spec["calculated"]]
    positions = {column["key"]: index + 1 for index, column in enumerate(all_columns)}
    for index, column in enumerate(all_columns, start=1):
        _literal(data.cell(3, index), column["header"])
        data.column_dimensions[get_column_letter(index)].width = min(max(len(column["header"]) + 5, 15), 28)
    _style_header(data, 3, len(all_columns))
    expected: dict[str, int | float] = {}
    formulas: dict[str, str] = {}
    for row_index, row in enumerate(spec["rows"], start=4):
        for column in spec["columns"]:
            cell = data.cell(row_index, positions[column["key"]])
            value = row[column["key"]]
            if column["type"] == "text":
                _literal(cell, value)
            else:
                cell.value = value
                cell.number_format = _number_format(column["type"])
            cell.font = Font(name=FONT_NAME, size=10, color=COLORS["ink"])
            cell.alignment = Alignment(vertical="center", horizontal="right" if column["type"] != "text" else "left")
        for column in spec["calculated"]:
            cell = data.cell(row_index, positions[column["key"]])
            left, right = (f"{get_column_letter(positions[key])}{row_index}" for key in column["inputs"])
            operator = {"add": "+", "subtract": "-", "multiply": "*", "divide": "/"}[column["operation"]]
            cell.value = f"={left}{operator}{right}"
            cell.number_format = column.get("numberFormat", "#,##0.00")
            cell.font = Font(name=FONT_NAME, size=10, color=COLORS["ink"])
            cell.alignment = Alignment(vertical="center", horizontal="right")
            formulas[f"{data.title}!{cell.coordinate}"] = cell.value
            expected[f"{data.title}!{cell.coordinate}"] = row[column["key"]]
        data.row_dimensions[row_index].height = 20
    last_row = len(spec["rows"]) + 3
    last_col = get_column_letter(len(all_columns))
    table = Table(displayName=spec["table"]["name"], ref=f"A3:{last_col}{last_row}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    data.add_table(table)
    data.print_area = f"A1:{last_col}{last_row}"
    for column in spec["columns"]:
        position = get_column_letter(positions[column["key"]])
        target = f"{position}4:{position}{last_row}"
        if column.get("choices"):
            choices = column["choices"]
            validation = DataValidation(type="list", formula1='"' + ",".join(choices) + '"', allow_blank=False)
            validation.error = "Choose a listed value"
            validation.showErrorMessage = True
            data.add_data_validation(validation)
            validation.add(target)
        if "minimum" in column or "maximum" in column:
            validation = DataValidation(
                type="whole" if column["type"] == "integer" else "decimal",
                operator="between",
                formula1=str(column.get("minimum", -10**15)),
                formula2=str(column.get("maximum", 10**15)),
                allow_blank=False,
            )
            validation.error = "Value is outside the permitted range"
            validation.showErrorMessage = True
            data.add_data_validation(validation)
            validation.add(target)
    for column in spec["calculated"]:
        if column.get("highlightNegative"):
            position = get_column_letter(positions[column["key"]])
            data.conditional_formatting.add(
                f"{position}4:{position}{last_row}",
                CellIsRule(operator="lessThan", formula=["0"], fill=PatternFill("solid", fgColor=COLORS["warning"])),
            )
    if summary_sheet:
        summary_sheet.sheet_view.showGridLines = False
        summary_sheet.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
        summary_sheet.page_setup.fitToWidth = 1
        summary_sheet.page_setup.fitToHeight = 1
        summary_sheet.page_setup.orientation = "landscape"
        summary_sheet.page_setup.paperSize = summary_sheet.PAPERSIZE_A4
        summary_sheet.page_margins = PageMargins(left=0.4, right=0.4, top=0.5, bottom=0.5, header=0.2, footer=0.2)
        _literal(summary_sheet["A1"], spec["title"])
        summary_sheet["A1"].font = Font(name=FONT_NAME, size=14, bold=True, color=COLORS["ink"])
        summary_sheet.column_dimensions["A"].width = 32
        summary_sheet.column_dimensions["B"].width = 20
        _literal(summary_sheet["A2"], "Показатель" if russian else "Metric")
        _literal(summary_sheet["B2"], "Значение" if russian else "Value")
        _style_header(summary_sheet, 2, 2)
        quoted_sheet = "'" + data.title.replace("'", "''") + "'"
        for row_index, item in enumerate(spec["summary"], start=3):
            _literal(summary_sheet.cell(row_index, 1), item["label"])
            summary_sheet.cell(row_index, 1).font = Font(name=FONT_NAME, size=10, color=COLORS["ink"])
            column_letter = get_column_letter(positions[item["column"]])
            source_range = f"{quoted_sheet}!{column_letter}4:{column_letter}{last_row}"
            function = "COUNTA" if item["operation"] == "count" else item["operation"].upper()
            cell = summary_sheet.cell(row_index, 2)
            cell.value = f"={function}({source_range})"
            cell.number_format = "#,##0" if function == "COUNTA" else "#,##0.00"
            cell.font = Font(name=FONT_NAME, size=10, bold=True, color=COLORS["ink"])
            formulas[f"Summary!{cell.coordinate}"] = cell.value
            values = [row[item["column"]] for row in spec["rows"]]
            expected[f"Summary!{cell.coordinate}"] = {
                "sum": lambda: sum(values),
                "average": lambda: sum(values) / len(values),
                "min": lambda: min(values),
                "max": lambda: max(values),
                "count": lambda: len(values),
            }[item["operation"]]()
        if spec["chart"]:
            chart_spec = spec["chart"]
            chart = BarChart() if chart_spec["kind"] == "bar" else LineChart()
            chart.title = chart_spec["title"]
            chart.style = 10
            chart.height = 10
            chart.width = 22
            chart.legend = None
            category_col = positions[chart_spec["category"]]
            value_col = positions[chart_spec["value"]]
            chart.add_data(Reference(data, min_col=value_col, min_row=3, max_row=last_row), titles_from_data=True)
            chart.set_categories(Reference(data, min_col=category_col, min_row=4, max_row=last_row))
            chart_row = max(7, len(spec["summary"]) + 4)
            summary_sheet.add_chart(chart, f"A{chart_row}")
        summary_sheet.print_area = (
            f"A1:L{chart_row + 22}" if spec["chart"] else f"A1:B{max(5, len(spec['summary']) + 3)}"
        )
    if len(spec["sources"]) > 1:
        sources_sheet = workbook.create_sheet("Sources")
        sources_sheet.sheet_view.showGridLines = False
        sources_sheet.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
        sources_sheet.page_setup.fitToWidth = 1
        sources_sheet.page_setup.fitToHeight = 1
        sources_sheet.page_setup.orientation = "landscape"
        sources_sheet.page_setup.paperSize = sources_sheet.PAPERSIZE_A4
        sources_sheet.page_margins = PageMargins(left=0.4, right=0.4, top=0.5, bottom=0.5, header=0.2, footer=0.2)
        sources_sheet.column_dimensions["A"].width = 32
        sources_sheet.column_dimensions["B"].width = 70
        _literal(sources_sheet["A1"], "Источники" if russian else "Sources")
        sources_sheet["A1"].font = Font(name=FONT_NAME, size=14, bold=True, color=COLORS["ink"])
        _literal(sources_sheet["A2"], "Источник" if russian else "Source")
        _literal(sources_sheet["B2"], "Адрес или файл" if russian else "Location")
        _style_header(sources_sheet, 2, 2)
        for row_index, source in enumerate(spec["sources"], start=3):
            _literal(sources_sheet.cell(row_index, 1), source["label"])
            _literal(sources_sheet.cell(row_index, 2), source.get("location") or source.get("url"))
            for cell in sources_sheet[row_index][:2]:
                cell.font = Font(name=FONT_NAME, size=10, color=COLORS["ink"])
        sources_sheet.print_area = f"A1:B{len(spec['sources']) + 2}"
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.calculation.calcMode = "auto"
    workbook.save(output)
    return formulas, expected, workbook.sheetnames


def _check_structure(output: Path, spec: dict[str, Any], formulas: dict[str, str], sheetnames: list[str]) -> list[str]:
    issues: list[str] = []
    workbook = load_workbook(output, data_only=False, keep_links=False)
    if workbook.sheetnames != sheetnames:
        issues.append("Workbook sheets changed after saving")
    data = workbook[spec["sheetName"]]
    if spec["table"]["name"] not in data.tables:
        issues.append("Native Excel Table is missing")
    if data.freeze_panes != "A4":
        issues.append("Data header freeze pane is missing")
    if spec["chart"] and len(workbook["Summary"]._charts) != 1:
        issues.append("Native chart is missing")
    for address, formula in formulas.items():
        sheet_name, cell_address = address.split("!", 1)
        if workbook[sheet_name][cell_address].value != formula:
            issues.append(f"Formula changed at {address}")
    for row in data.iter_rows(min_row=4, max_row=len(spec["rows"]) + 3):
        for cell in row[: len(spec["columns"])]:
            if cell.data_type == "f":
                issues.append(f"A source value became a formula at {data.title}!{cell.coordinate}")
    workbook.close()
    return issues


def _office_convert(source: Path, outdir: Path, format_name: str, profile: Path) -> Path:
    soffice = shutil.which("soffice")
    if not soffice:
        raise RuntimeError("LibreOffice is unavailable")
    outdir.mkdir(parents=True, exist_ok=True)
    command = [
        soffice,
        f"-env:UserInstallation={profile.as_uri()}",
        "--headless",
        "--convert-to",
        format_name,
        "--outdir",
        str(outdir),
        str(source),
    ]
    result = subprocess.run(command, capture_output=True, text=True, timeout=120, check=False)
    target = outdir / f"{source.stem}.{format_name}"
    if result.returncode != 0 or not target.is_file() or target.stat().st_size == 0:
        detail = result.stderr.strip() or result.stdout.strip() or "no output file"
        raise RuntimeError(f"LibreOffice {format_name} conversion failed (exit {result.returncode}): {detail}")
    return target


def _check_recalculated(path: Path, expected: dict[str, int | float]) -> list[str]:
    issues: list[str] = []
    workbook = load_workbook(path, data_only=True, keep_links=False)
    for sheet in workbook:
        for row in sheet.iter_rows():
            for cell in row:
                if cell.data_type == "e":
                    issues.append(f"Formula error at {sheet.title}!{cell.coordinate}: {cell.value}")
    for address, expected_value in expected.items():
        sheet_name, cell_address = address.split("!", 1)
        actual = workbook[sheet_name][cell_address].value
        if isinstance(actual, bool) or not isinstance(actual, (int, float)) or not math.isclose(
            actual, expected_value, rel_tol=1e-9, abs_tol=1e-6
        ):
            issues.append(f"Recalculated value differs at {address}: {actual!r} versus {expected_value!r}")
    workbook.close()
    return issues


def _check_render(pdf: Path, sheetnames: list[str], scratch: Path, exact_pages: bool) -> list[str]:
    rasterizer = shutil.which("pdftoppm")
    if not rasterizer:
        return ["Poppler rasterizer is unavailable"]
    result = subprocess.run(
        [rasterizer, "-r", "90", "-png", str(pdf), str(scratch / "page")],
        capture_output=True,
        text=True,
        timeout=120,
        check=False,
    )
    images = sorted(scratch.glob("page-*.png"))
    if result.returncode != 0 or len(images) < len(sheetnames):
        return ["Not every worksheet rendered to a PDF page"]
    issues: list[str] = []
    if exact_pages and len(images) != len(sheetnames):
        issues.append(f"Small workbook rendered {len(images)} pages for {len(sheetnames)} sheets")
    for index, image_path in enumerate(images, start=1):
        with Image.open(image_path) as image:
            page = image.convert("RGB")
            difference = ImageChops.difference(page, Image.new("RGB", page.size, "white"))
            if difference.getbbox() is None:
                issues.append(f"Rendered page {index} is blank")
    return issues


def _check(name: str, issues: list[str], success: str) -> dict[str, str]:
    return {"name": name, "status": "failed" if issues else "passed", "message": "; ".join(issues) if issues else success}


def _issue(code: str, message: str) -> dict[str, str]:
    return {"code": code, "severity": "critical", "message": message}


def _report(spec: dict[str, Any], checks: list[dict[str, str]], issues: list[dict[str, str]], preview: Path | None) -> dict[str, Any]:
    return {
        "status": "needs_review" if issues else "ready",
        "format": "xlsx",
        "sourceFileIds": list(spec["job"]["sourceFileIds"]),
        "previewAssets": ([{"filename": preview.name, "kind": "pdf", "delivery": "preview_only"}] if preview else []),
        "qaChecks": checks,
        "issues": issues,
        "changeLog": list(spec["changes"]) or [{"target": spec["job"]["filename"], "summary": "Created an editable formula-driven workbook"}],
        "skillVersion": SKILL_VERSION,
        "repairIterations": spec["repairIterations"],
    }


def main() -> int:
    if len(sys.argv) != 3:
        print("usage: build_spreadsheet.py SPEC.json OUTPUT.xlsx", file=sys.stderr)
        return 2
    spec_path = Path(sys.argv[1]).resolve()
    output = Path(sys.argv[2]).resolve()
    if spec_path == output or spec_path == Path(f"{output}.artifact-report.json").resolve():
        raise SpecError("Specification and output paths must differ")
    raw = json.loads(spec_path.read_text(encoding="utf-8"))
    spec = _validate(raw, output)
    output.parent.mkdir(parents=True, exist_ok=True)
    report_path = Path(f"{output}.artifact-report.json")
    preview_path = output.with_name(f"_qa_{output.stem}-preview.pdf")
    if any(path.exists() for path in (output, report_path, preview_path)):
        raise SpecError("Existing output cannot be overwritten; use a new versioned filename")
    formulas, expected, sheetnames = _build(spec, output)
    checks: list[dict[str, str]] = []
    issues: list[dict[str, str]] = []
    structure_issues = _check_structure(output, spec, formulas, sheetnames)
    checks.append(_check("structure-and-formulas", structure_issues, "Workbook reopens with native table, chart, and editable formulas"))
    issues.extend(_issue("structure-and-formulas", problem) for problem in structure_issues)
    with tempfile.TemporaryDirectory(prefix="xlsx-qa-") as folder:
        scratch = Path(folder)
        try:
            recalculated = _office_convert(output, scratch / "recalculated", "xlsx", scratch / "profile-recalc")
            value_issues = _check_recalculated(recalculated, expected)
        except (OSError, RuntimeError, subprocess.TimeoutExpired) as exc:
            value_issues = [str(exc)]
        checks.append(_check("recalculated-values", value_issues, "LibreOffice values match independent calculations"))
        issues.extend(_issue("recalculated-values", problem) for problem in value_issues)
        try:
            pdf = _office_convert(output, scratch / "rendered", "pdf", scratch / "profile-render")
            compact = len(spec["rows"]) <= 20 and len(spec["columns"]) + len(spec["calculated"]) <= 8
            render_issues = _check_render(pdf, sheetnames, scratch, exact_pages=compact)
            shutil.copyfile(pdf, preview_path)
        except (OSError, RuntimeError, subprocess.TimeoutExpired) as exc:
            render_issues = [str(exc)]
        checks.append(_check("render", render_issues, "Every worksheet rendered to nonblank PDF pages"))
        issues.extend(_issue("render", problem) for problem in render_issues)
    report = _report(spec, checks, issues, preview_path if preview_path.is_file() else None)
    report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"artifact": str(output), "report": str(report_path), "status": report["status"]}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
